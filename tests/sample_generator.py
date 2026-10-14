from __future__ import annotations

import io
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

from core.parser import NAMESPACES

SAMPLE_TIMING_XML = """
<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
          xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:tnLst>
    <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
      <p:childTnLst>
        <p:seq concurrent="1" nextAc="seek">
          <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
            <p:childTnLst>
              <!-- 1. Entrance Fade on Shape 2 -->
              <p:par>
                <p:cTn id="3" fill="hold">
                  <p:stCondLst>
                    <p:cond delay="0"/>
                  </p:stCondLst>
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="4" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="clickEffect">
                        <p:childTnLst>
                          <p:set>
                            <p:cBhvr>
                              <p:cTn id="5" dur="1" fill="hold"/>
                              <p:tgtEl>
                                <p:spTgt spid="2"/>
                              </p:tgtEl>
                              <p:attrNameLst>
                                <p:attrName>style.visibility</p:attrName>
                              </p:attrNameLst>
                            </p:cBhvr>
                            <p:to>
                              <p:strVal val="visible"/>
                            </p:to>
                          </p:set>
                          <p:animEffect transition="in" filter="fade">
                            <p:cBhvr>
                              <p:cTn id="6" dur="500"/>
                              <p:tgtEl>
                                <p:spTgt spid="2"/>
                              </p:tgtEl>
                            </p:cBhvr>
                          </p:animEffect>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
              <!-- 2. Emphasis Spin on Shape 3 -->
              <p:par>
                <p:cTn id="7" fill="hold">
                  <p:stCondLst>
                    <p:cond delay="0"/>
                  </p:stCondLst>
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="8" presetID="23" presetClass="emph" presetSubtype="0" fill="hold" nodeType="clickEffect">
                        <p:childTnLst>
                          <p:animRot by="21600000">
                            <p:cBhvr>
                              <p:cTn id="9" dur="1000"/>
                              <p:tgtEl>
                                <p:spTgt spid="3"/>
                              </p:tgtEl>
                            </p:cBhvr>
                          </p:animRot>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
              <!-- 3. Exit Fade on Shape 4 -->
              <p:par>
                <p:cTn id="10" fill="hold">
                  <p:stCondLst>
                    <p:cond delay="200"/>
                  </p:stCondLst>
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="11" presetID="10" presetClass="exit" presetSubtype="0" fill="hold" nodeType="afterEffect">
                        <p:childTnLst>
                          <p:animEffect transition="out" filter="fade">
                            <p:cBhvr>
                              <p:cTn id="12" dur="500"/>
                              <p:tgtEl>
                                <p:spTgt spid="4"/>
                              </p:tgtEl>
                            </p:cBhvr>
                          </p:animEffect>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
              <!-- 4. Motion Path on Shape 5 -->
              <p:par>
                <p:cTn id="13" fill="hold">
                  <p:stCondLst>
                    <p:cond delay="0"/>
                  </p:stCondLst>
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="14" presetID="55" presetClass="path" presetSubtype="0" fill="hold" nodeType="clickEffect">
                        <p:childTnLst>
                          <p:animMotion origin="layout" path="M 0 0 L 0.25 0.25 E" ptsTypes="">
                            <p:cBhvr>
                              <p:cTn id="15" dur="1500"/>
                              <p:tgtEl>
                                <p:spTgt spid="5"/>
                              </p:tgtEl>
                            </p:cBhvr>
                          </p:animMotion>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
              <!-- 5. By-Paragraph Build on Shape 6 -->
              <p:par>
                <p:cTn id="16" fill="hold">
                  <p:stCondLst>
                    <p:cond delay="0"/>
                  </p:stCondLst>
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="17" presetID="1" presetClass="entr" fill="hold" nodeType="clickEffect">
                        <p:childTnLst>
                          <p:set>
                            <p:cBhvr>
                              <p:cTn id="18" dur="1"/>
                              <p:tgtEl>
                                <p:spTgt spid="6">
                                  <p:txEl>
                                    <p:pRg st="0" end="0"/>
                                  </p:txEl>
                                </p:spTgt>
                              </p:tgtEl>
                            </p:cBhvr>
                          </p:set>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
              <!-- 6. Target on Shape 7 (will be DELETED in remediated deck) -->
              <p:par>
                <p:cTn id="19" fill="hold">
                  <p:stCondLst>
                    <p:cond delay="0"/>
                  </p:stCondLst>
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="20" presetID="10" presetClass="entr" fill="hold" nodeType="clickEffect">
                        <p:childTnLst>
                          <p:animEffect transition="in" filter="fade">
                            <p:cBhvr>
                              <p:cTn id="21" dur="500"/>
                              <p:tgtEl>
                                <p:spTgt spid="7"/>
                              </p:tgtEl>
                            </p:cBhvr>
                          </p:animEffect>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
            </p:childTnLst>
          </p:cTn>
        </p:seq>
        <!-- 7. Interactive Sequence Trigger on Shape 2 -->
        <p:seq concurrent="1" nextAc="seek">
          <p:cTn id="22" dur="indefinite" nodeType="interactiveSeq">
            <p:stCondLst>
              <p:cond evt="onClick" delay="0">
                <p:tgtEl>
                  <p:spTgt spid="2"/>
                </p:tgtEl>
              </p:cond>
            </p:stCondLst>
            <p:childTnLst>
              <p:par>
                <p:cTn id="23" presetID="23" presetClass="emph" fill="hold" nodeType="clickEffect">
                  <p:childTnLst>
                    <p:animRot by="10800000">
                      <p:cBhvr>
                        <p:cTn id="24" dur="500"/>
                        <p:tgtEl>
                          <p:spTgt spid="3"/>
                        </p:tgtEl>
                      </p:cBhvr>
                    </p:animRot>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
            </p:childTnLst>
          </p:cTn>
        </p:seq>
      </p:childTnLst>
    </p:cTn>
  </p:tnLst>
  <p:bldLst>
    <p:bldP spid="6" grpId="0" build="p"/>
    <p:bldP spid="7" grpId="0" build="p"/>
  </p:bldLst>
</p:timing>
"""

SAMPLE_TRANSITION_XML = """
<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">
  <p:wipe dir="r"/>
</p:transition>
"""


def create_sample_decks() -> tuple[bytes, bytes]:
    """
    Create:
    1. Original deck with shapes (ids 2, 3, 4, 5, 6, 7), transition, and full timing tree.
    2. Remediated deck where:
       - Shape ids are changed (e.g. 2->202, 3->303, 4->404, 5->505, 6->606)
       - Shape 3 is renamed
       - Shape 7 is deleted
       - No initial animations or transitions
    """
    # -------------------------------------------------------------------------
    # 1. Create Original Deck
    # -------------------------------------------------------------------------
    orig_prs = Presentation()
    blank_layout = orig_prs.slide_layouts[6]
    orig_slide = orig_prs.slides.add_slide(blank_layout)

    # Add shapes:
    # Shape 2: Title box
    s2 = orig_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    s2.text = "Original Title Box"
    s2.name = "Title Box"

    # Shape 3: Spinner rectangle
    s3 = orig_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.5), Inches(2), Inches(1.5))
    s3.text = "Spinner Box"
    s3.name = "Spinner Box"

    # Shape 4: Exit Oval
    s4 = orig_slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(2.5), Inches(2), Inches(1.5))
    s4.text = "Exit Oval"
    s4.name = "Exit Oval"

    # Shape 5: Motion Path Star
    s5 = orig_slide.shapes.add_shape(getattr(MSO_SHAPE, "FIVE_POINT_STAR", MSO_SHAPE.ROUNDED_RECTANGLE), Inches(1), Inches(4.5), Inches(2), Inches(1.5))
    s5.text = "Motion Star"
    s5.name = "Motion Star"

    # Shape 6: Paragraph bullet box
    s6 = orig_slide.shapes.add_textbox(Inches(4), Inches(4.5), Inches(4), Inches(2))
    s6.name = "Bullet Box"
    tf = s6.text_frame
    p0 = tf.paragraphs[0]
    p0.text = "First paragraph line"
    p1 = tf.add_paragraph()
    p1.text = "Second paragraph line"

    # Shape 7: Shape to be deleted in remediation
    s7 = orig_slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(7), Inches(1), Inches(1.5), Inches(1.5))
    s7.text = "Deleted Diamond"
    s7.name = "Deleted Diamond"

    # Inject explicit IDs: 2, 3, 4, 5, 6, 7
    shape_id_pairs = [(s2, "2"), (s3, "3"), (s4, "4"), (s5, "5"), (s6, "6"), (s7, "7")]
    for shp, explicit_id in shape_id_pairs:
        cNvPr = shp._element.find(f".//{{{NAMESPACES['p']}}}cNvPr")
        if cNvPr is not None:
            cNvPr.attrib["id"] = explicit_id

    # Inject Transition & Timing XML into slide XML
    sld_elem = orig_slide._element
    trans_elem = etree.fromstring(SAMPLE_TRANSITION_XML)
    timing_elem = etree.fromstring(SAMPLE_TIMING_XML)

    sld_elem.append(trans_elem)
    sld_elem.append(timing_elem)

    orig_bio = io.BytesIO()
    orig_prs.save(orig_bio)
    orig_bytes = orig_bio.getvalue()

    # -------------------------------------------------------------------------
    # 2. Create Remediated Deck
    # -------------------------------------------------------------------------
    rem_prs = Presentation()
    rem_slide = rem_prs.slides.add_slide(blank_layout)

    # Remediated Shape 202 (corresponds to 2)
    rs2 = rem_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    rs2.text = "Original Title Box"
    rs2.name = "Title Box"

    # Remediated Shape 303 (corresponds to 3, but renamed and with changed ID)
    rs3 = rem_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2.5), Inches(2), Inches(1.5))
    rs3.text = "Spinner Box"
    rs3.name = "Renamed Spinner Box"

    # Remediated Shape 404 (corresponds to 4)
    rs4 = rem_slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(2.5), Inches(2), Inches(1.5))
    rs4.text = "Exit Oval"
    rs4.name = "Exit Oval"

    # Remediated Shape 505 (corresponds to 5)
    rs5 = rem_slide.shapes.add_shape(getattr(MSO_SHAPE, "FIVE_POINT_STAR", MSO_SHAPE.ROUNDED_RECTANGLE), Inches(1), Inches(4.5), Inches(2), Inches(1.5))
    rs5.text = "Motion Star"
    rs5.name = "Motion Star"

    # Remediated Shape 606 (corresponds to 6)
    rs6 = rem_slide.shapes.add_textbox(Inches(4), Inches(4.5), Inches(4), Inches(2))
    rs6.name = "Bullet Box"
    rtf = rs6.text_frame
    rp0 = rtf.paragraphs[0]
    rp0.text = "First paragraph line"
    rp1 = rtf.add_paragraph()
    rp1.text = "Second paragraph line"

    # Notice: Shape 7 is NOT added to rem_slide! It is deleted.

    # Assign new IDs: 202, 303, 404, 505, 606
    rem_id_pairs = [(rs2, "202"), (rs3, "303"), (rs4, "404"), (rs5, "505"), (rs6, "606")]
    for shp, explicit_id in rem_id_pairs:
        cNvPr = shp._element.find(f".//{{{NAMESPACES['p']}}}cNvPr")
        if cNvPr is not None:
            cNvPr.attrib["id"] = explicit_id

    rem_bio = io.BytesIO()
    rem_prs.save(rem_bio)
    rem_bytes = rem_bio.getvalue()

    return orig_bytes, rem_bytes
